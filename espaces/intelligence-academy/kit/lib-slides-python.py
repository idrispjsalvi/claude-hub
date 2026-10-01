#!/usr/bin/env python3
"""
Boîte à outils partagée - slides Formation « Piloter Claude Code comme un power-user ».
Charte TIA (Bricolage Grotesque, teal/gold). Importée par generate_formation_cc_m*.py.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SLIDES_DIR = os.path.dirname(SCRIPT_DIR)
BASE = os.path.dirname(SLIDES_DIR)


def _trouver_logos_dir():
    """Le dossier qui contient VRAIMENT les logos.

    ⛔ `docs/formations/logos/` ne contient que 6 SVG d'outils : ni `ia-round.png`, ni `claude.png`.
    Comme chaque `add_picture` est gardé par un `os.path.exists`, le logo était simplement SAUTÉ —
    les decks sortaient sans marque, sans que rien ne le dise (constaté le 02/09/2026).

    Ordre de recherche : `TIA_LOGOS_DIR` (posé par qui veut décider) · un `logos/` à côté du script
    (c'est la forme du kit envoyé aux formateurs) · `docs/formations/logos` · `docs/brochures/assets`
    (où ils vivent réellement dans le dépôt). Le premier qui porte `ia-round.png` gagne.
    """
    candidats = [
        os.environ.get("TIA_LOGOS_DIR"),
        os.path.join(SCRIPT_DIR, "logos"),
        os.path.join(BASE, "logos"),
        os.path.join(os.path.dirname(BASE), "brochures", "assets"),
    ]
    for c in candidats:
        if c and os.path.exists(os.path.join(c, "ia-round.png")):
            return c
    return os.path.join(BASE, "logos")


LOGOS_DIR = _trouver_logos_dir()
IA_LOGO = os.path.join(LOGOS_DIR, "ia-round.png")

C_HEADER_LEFT  = RGBColor(0x0A, 0x25, 0x30)
C_HEADER_RIGHT = RGBColor(0x1E, 0x55, 0x65)
C_GOLD         = RGBColor(0xE8, 0xA0, 0x00)
C_WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK        = RGBColor(0x00, 0x00, 0x00)
C_TITLE        = RGBColor(0x1A, 0x1A, 0x1A)
C_BODY         = RGBColor(0x33, 0x33, 0x33)
C_TEAL         = RGBColor(0x1B, 0x5E, 0x6B)
C_TEAL_LIGHT   = RGBColor(0x5A, 0x8A, 0x8F)
C_DARK_BLUE    = RGBColor(0x0F, 0x2B, 0x3A)
C_GRAY_BG      = RGBColor(0xED, 0xED, 0xED)

W = Inches(13.333); H = Inches(7.5)
FONT = FONT_BODY = "Bricolage Grotesque"
MONO = "Consolas"


def new_prs():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H; return prs

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def add_rect(slide, left, top, width, height, fill_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    return shape

def add_rounded_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    shape = slide.shapes.add_shape(5, left, top, width, height)
    if fill_color:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color; shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_textbox(slide, left, top, width, height, text, font_name=FONT,
                font_size=18, bold=False, color=C_BODY, align=PP_ALIGN.LEFT, word_wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame; tf.word_wrap = word_wrap
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = font_name; run.font.size = Pt(font_size)
    run.font.bold = bold; run.font.color.rgb = color
    return txBox, tf

def _apply_gradient(shape, color_left, color_right):
    ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    spPr = shape._element.spPr
    for tag in ('solidFill', 'gradFill', 'noFill', 'pattFill'):
        for el in spPr.findall(f'{{{ns}}}{tag}'):
            spPr.remove(el)
    gradFill = etree.SubElement(spPr, f'{{{ns}}}gradFill')
    gsLst = etree.SubElement(gradFill, f'{{{ns}}}gsLst')
    gs1 = etree.SubElement(gsLst, f'{{{ns}}}gs', attrib={'pos': '0'})
    etree.SubElement(gs1, f'{{{ns}}}srgbClr', attrib={'val': str(color_left)})
    gs2 = etree.SubElement(gsLst, f'{{{ns}}}gs', attrib={'pos': '100000'})
    etree.SubElement(gs2, f'{{{ns}}}srgbClr', attrib={'val': str(color_right)})
    etree.SubElement(gradFill, f'{{{ns}}}lin', attrib={'ang': '0', 'scaled': '1'})

def add_header(slide, title_text):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), W, Inches(1.25))
    shape.line.fill.background()
    _apply_gradient(shape, C_HEADER_LEFT, C_HEADER_RIGHT)
    add_rect(slide, Inches(0), Inches(1.25), Inches(13.333), Inches(0.02), C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(0.2), Inches(11.5), Inches(0.9),
                title_text, font_size=34, bold=True, color=C_WHITE)

def add_subtitle(slide, text):
    add_textbox(slide, Inches(0.7), Inches(1.6), Inches(11.5), Inches(0.7),
                text, font_size=26, bold=True, color=C_TITLE)

def add_brand_footer(slide, y=Inches(6.8), centered=False):
    logo_size = Inches(0.35)
    if centered:
        logo_x, text_x = Inches(5.3), Inches(5.75)
    else:
        logo_x, text_x = Inches(0.5), Inches(0.95)
    if os.path.exists(IA_LOGO):
        try: slide.shapes.add_picture(IA_LOGO, logo_x, y, logo_size, logo_size)
        except: pass
    add_textbox(slide, text_x, y - Inches(0.02), Inches(4.0), Inches(0.4),
                "intelligence academy", font_size=14, bold=False, color=C_WHITE)

def _split_into_blocks(lines):
    blocks, current = [], []
    for line in lines:
        if line == "":
            if current: blocks.append(current); current = []
        else:
            current.append(line)
    if current: blocks.append(current)
    return blocks

def _draw_blocks_on_slide(slide, top, blocks, visible_block_count, font_size=20, width=Inches(11.9)):
    """Markers: '**' bold teal, '> ' gold quote, '` ' mono code."""
    y = top; line_h = Inches(0.45); gap = Inches(0.12); block_gap = Inches(0.26)
    for b_idx, block in enumerate(blocks):
        if b_idx >= visible_block_count: break
        if b_idx > 0: y += block_gap
        for line in block:
            color = C_BODY; bold = False; fname = FONT_BODY
            if line.startswith("**"):
                line = line[2:]; bold = True; color = C_TEAL
            elif line.startswith("> "):
                line = "« " + line[2:] + " »"; color = C_GOLD; bold = True
            elif line.startswith("` "):
                line = line[2:]; fname = MONO; color = C_DARK_BLUE
            tb = slide.shapes.add_textbox(Inches(0.7), y, width, line_h)
            tf = tb.text_frame; tf.word_wrap = True
            run = tf.paragraphs[0].add_run(); run.text = line
            run.font.name = fname; run.font.size = Pt(font_size); run.font.bold = bold
            run.font.color.rgb = color
            y += line_h + gap


def slide_cover(prs, kicker, title, baseline, subtitle):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), W, H, C_HEADER_LEFT)
    add_rect(slide, Inches(0), Inches(4.6), W, Inches(2.9), C_HEADER_RIGHT)
    add_rect(slide, Inches(0.7), Inches(1.05), Inches(2.4), Inches(0.04), C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(1.2), Inches(11.5), Inches(0.8),
                kicker, font_size=18, bold=True, color=C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(2.1), Inches(11.9), Inches(1.5),
                title, font_size=50, bold=True, color=C_WHITE)
    add_textbox(slide, Inches(0.7), Inches(3.7), Inches(11.5), Inches(0.8),
                baseline, font_name=FONT_BODY, font_size=26, bold=False, color=C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(4.9), Inches(11.5), Inches(0.8),
                subtitle, font_name=FONT_BODY, font_size=20, bold=False, color=C_TEAL_LIGHT)
    add_brand_footer(slide, y=Inches(6.8))

def slide_section(prs, num, title, subtitle=None):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), W, H, C_DARK_BLUE)
    add_rect(slide, Inches(0.7), Inches(2.15), Inches(3.0), Inches(0.04), C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(1.5), Inches(11.5), Inches(0.8),
                f"Séquence {num}", font_size=22, bold=True, color=C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(2.5), Inches(11.7), Inches(1.8),
                title, font_size=42, bold=True, color=C_WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.7), Inches(4.4), Inches(11.5), Inches(0.8),
                    subtitle, font_name=FONT_BODY, font_size=20, bold=False, color=C_TEAL_LIGHT)
    add_brand_footer(slide, y=Inches(6.8))

def slide_content(prs, header_title, subtitle, bullets, font_size=20):
    blocks = _split_into_blocks(bullets)
    top = Inches(2.55) if subtitle else Inches(1.7)
    out = []
    for i in range(1, len(blocks) + 1):
        slide = blank_slide(prs)
        add_header(slide, header_title)
        if subtitle: add_subtitle(slide, subtitle)
        _draw_blocks_on_slide(slide, top, blocks, i, font_size)
        out.append(slide)
    return out

def slide_pipeline(prs, header_title, subtitle, steps, note=None):
    """Diagramme horizontal de N étapes (wrap à 4 par ligne)."""
    slide = blank_slide(prs)
    add_header(slide, header_title)
    if subtitle: add_subtitle(slide, subtitle)
    w = Inches(2.55); h = Inches(0.85); per_row = 4
    start_x = Inches(0.7); start_y = Inches(3.0)
    for i, s in enumerate(steps):
        col = i % per_row; row = i // per_row
        cx = start_x + col * (w + Inches(0.42))
        cy = start_y + row * (h + Inches(0.7))
        add_rounded_rect(slide, cx, cy, w, h, fill_color=C_TEAL if i % 2 == 0 else C_HEADER_RIGHT)
        tb = slide.shapes.add_textbox(cx + Inches(0.1), cy + Inches(0.18), w - Inches(0.2), Inches(0.55))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = s
        run.font.name = FONT; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = C_WHITE
        if col < per_row - 1 and i < len(steps) - 1:
            add_textbox(slide, cx + w + Inches(0.04), cy + Inches(0.2), Inches(0.34), Inches(0.5), "→",
                        font_size=22, bold=True, color=C_GOLD, align=PP_ALIGN.CENTER)
    if note:
        add_textbox(slide, Inches(0.7), Inches(5.7), Inches(12.0), Inches(0.7),
                    note, font_size=18, bold=True, color=C_TEAL)
    return slide

def slide_recap(prs, title, items):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), W, H, C_DARK_BLUE)
    add_textbox(slide, Inches(0.7), Inches(1.0), Inches(11.5), Inches(0.8),
                "À retenir", font_size=36, bold=True, color=C_GOLD)
    add_textbox(slide, Inches(0.7), Inches(1.8), Inches(11.5), Inches(0.6),
                title, font_name=FONT_BODY, font_size=18, bold=False, color=C_TEAL_LIGHT)
    add_rect(slide, Inches(0.7), Inches(2.5), Inches(3.0), Inches(0.03), C_GOLD)
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(11.9), Inches(3.7))
    tf = tb.text_frame; tf.word_wrap = True; first = True
    for item in items:
        if first: p = tf.paragraphs[0]; first = False
        else: p = tf.add_paragraph()
        p.space_before = Pt(12)
        run = p.add_run(); run.text = item
        run.font.name = FONT_BODY; run.font.size = Pt(20); run.font.color.rgb = C_WHITE
    add_brand_footer(slide, y=Inches(6.95))

def slide_questions(prs):
    slide = blank_slide(prs)
    add_rect(slide, Inches(0), Inches(0), W, H, C_BLACK)
    add_textbox(slide, Inches(0), Inches(2.8), W, Inches(2.0),
                "Des questions ?", font_size=64, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_brand_footer(slide, y=Inches(5.2), centered=True)

def save(prs, output_path):
    prs.save(output_path)
    n = len(prs.slides._sldIdLst)
    print(f"OK -> {output_path} ({n} slides)")
